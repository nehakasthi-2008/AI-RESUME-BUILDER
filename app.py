import io
import os
import re
import streamlit as st
import ollama

# Optional libraries (add them to requirements.txt)
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None
try:
    from docx import Document
except ImportError:
    Document = None
try:
    from fpdf import FPDF
except ImportError:
    FPDF = None

st.set_page_config(
    page_title="AI Resume Builder",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded",
)


def rerun():
    """Works on old and new Streamlit versions."""
    fn = getattr(st, "rerun", None) or getattr(st, "experimental_rerun")
    fn()


# ---------- OPTIONAL: your own AI image (direct link ending in .jpg/.png/.webp) ----------
CUSTOM_IMAGE_URL = ""   # e.g. "https://your-site.com/my-ai-background.jpg"

# ---------- OLLAMA CONNECTION ----------
try:
    OLLAMA_HOST = st.secrets["OLLAMA_HOST"]
except Exception:
    OLLAMA_HOST = os.environ.get("OLLAMA_HOST", "http://localhost:11434")

client = ollama.Client(host=OLLAMA_HOST)

# ---------- EDUCATION OPTIONS ----------
PH = "-- Select --"
YEARS = [PH] + [str(y) for y in range(2032, 2004, -1)]
BOARDS_10 = [PH, "State Board", "CBSE", "ICSE", "Other"]
BOARDS_12 = [PH, "State Board (Intermediate)", "CBSE", "ICSE", "Other"]
STREAMS = [PH, "MPC", "BiPC", "MEC", "CEC", "HEC", "Diploma / Polytechnic", "Other"]
DEGREES = [PH, "B.Tech", "B.E", "B.Sc", "BCA", "B.Com", "BBA", "Diploma",
           "M.Tech", "MCA", "MBA", "M.Sc", "Other"]
BRANCHES = [PH, "CSE", "AIML", "EC", "ECE", "DS", "IT", "CS", "EEE",
            "Mechanical", "Civil", "Other"]

# ---------- COLOUR THEMES ----------
THEMES = {
    "⚡ Neon Cyber": dict(
        bg1="#050816", bg2="#12073a", c1="#00e5ff", c2="#c026d3", text="#ffffff",
        card="rgba(255,255,255,0.07)", input="rgba(0,0,0,0.45)", panel="rgba(5,8,22,0.88)",
        solid="#0b1030", line="rgba(0,229,255,0.12)", border="rgba(0,229,255,0.35)", btn="#050816",
    ),
    "🌌 Aurora": dict(
        bg1="#06141b", bg2="#052e2b", c1="#00ffa3", c2="#0ea5e9", text="#ffffff",
        card="rgba(255,255,255,0.07)", input="rgba(0,0,0,0.45)", panel="rgba(6,20,27,0.88)",
        solid="#0a1f27", line="rgba(0,255,163,0.12)", border="rgba(0,255,163,0.35)", btn="#04121a",
    ),
    "🔥 Ember": dict(
        bg1="#1a0a0a", bg2="#3b0d0d", c1="#ff5722", c2="#fbbf24", text="#ffffff",
        card="rgba(255,255,255,0.07)", input="rgba(0,0,0,0.45)", panel="rgba(26,10,10,0.88)",
        solid="#2a1010", line="rgba(255,87,34,0.14)", border="rgba(255,152,0,0.4)", btn="#1a0a0a",
    ),
    "🍃 Forest": dict(
        bg1="#04150e", bg2="#0a2e1c", c1="#34d399", c2="#a3e635", text="#ffffff",
        card="rgba(255,255,255,0.07)", input="rgba(0,0,0,0.45)", panel="rgba(4,21,14,0.88)",
        solid="#08241a", line="rgba(52,211,153,0.13)", border="rgba(163,230,53,0.35)", btn="#04150e",
    ),
    "🌸 Blush Light": dict(
        bg1="#fff1f7", bg2="#f3e8ff", c1="#ec4899", c2="#8b5cf6", text="#2d1b3d",
        card="rgba(255,255,255,0.7)", input="rgba(255,255,255,0.95)", panel="rgba(255,255,255,0.88)",
        solid="#ffffff", line="rgba(236,72,153,0.14)", border="rgba(139,92,246,0.35)", btn="#ffffff",
    ),
    "🌤️ Clean Light": dict(
        bg1="#eef4ff", bg2="#e0f2fe", c1="#2563eb", c2="#06b6d4", text="#0f172a",
        card="rgba(255,255,255,0.75)", input="rgba(255,255,255,0.95)", panel="rgba(255,255,255,0.9)",
        solid="#ffffff", line="rgba(37,99,235,0.13)", border="rgba(37,99,235,0.35)", btn="#ffffff",
    ),
}

BG_STYLES = ["Animated grid", "Aurora glow", "Floating resumes", "Soft dots"]

# ---------- SESSION STATE ----------
if "resume_text" not in st.session_state:
    st.session_state.resume_text = ""
if "history" not in st.session_state:
    st.session_state.history = []
if "current" not in st.session_state:
    st.session_state.current = None

# ---------- SIDEBAR: THEME + SETTINGS ----------
with st.sidebar:
    st.header("🎨 Theme")
    theme_name = st.selectbox("Colour theme", list(THEMES.keys()))
    bg_style = st.selectbox("Background style", BG_STYLES)

    st.header("⚙️ Settings")
    model_name = st.text_input("Ollama model", value="llama3.2")
    st.caption(f"Server: {OLLAMA_HOST}")

t = THEMES[theme_name]

# ---------- CSS ----------
root_css = (
    ":root{"
    f"--bg1:{t['bg1']};--bg2:{t['bg2']};--c1:{t['c1']};--c2:{t['c2']};"
    f"--text:{t['text']};--card:{t['card']};--input:{t['input']};--panel:{t['panel']};"
    f"--solid:{t['solid']};--line:{t['line']};--border:{t['border']};--btn:{t['btn']};"
    "}"
)

main_css = """
/* Base background */
.stApp {
    background: linear-gradient(135deg, var(--bg1), var(--bg2));
    color: var(--text);
}
[data-testid="stAppViewContainer"],
[data-testid="stMain"],
header[data-testid="stHeader"] {
    background: transparent !important;
}
[data-testid="stMain"] {
    position: relative;
    z-index: 1;
}

/* Main content card */
.block-container {
    max-width: 1100px;
    background: var(--card);
    border: 1px solid var(--border);
    border-radius: 20px;
    padding: 2rem 2.5rem !important;
    margin-top: 1.2rem;
    margin-bottom: 1.2rem;
    backdrop-filter: blur(10px);
    box-shadow: 0 6px 24px rgba(0, 0, 0, 0.25);
}

/* Background layers (behind the content) */
.bg-photo {
    position: fixed; inset: 0; z-index: -1; pointer-events: none;
    background-size: cover; background-position: center; opacity: 0.35;
}
.bg-grid {
    position: fixed; inset: 0; z-index: -1; pointer-events: none;
    background-image:
        linear-gradient(var(--line) 1px, transparent 1px),
        linear-gradient(90deg, var(--line) 1px, transparent 1px);
    background-size: 50px 50px;
    animation: gridmove 6s linear infinite;
}
@keyframes gridmove { to { background-position: 50px 50px; } }
.bg-dots {
    position: fixed; inset: 0; z-index: -1; pointer-events: none;
    background-image: radial-gradient(var(--border) 1.6px, transparent 1.7px);
    background-size: 30px 30px;
    animation: gridmove 10s linear infinite;
}
.orb {
    position: fixed; border-radius: 50%; filter: blur(85px);
    opacity: 0.5; z-index: -1; pointer-events: none;
    animation: drift 14s ease-in-out infinite alternate;
}
.o1 { width: 420px; height: 420px; background: var(--c1); top: -120px; left: -100px; }
.o2 { width: 460px; height: 460px; background: var(--c2); bottom: -140px; right: -120px; animation-delay: 3s; }
@keyframes drift {
    from { transform: translate(0, 0) scale(1); }
    to   { transform: translate(120px, 80px) scale(1.2); }
}
.blob {
    position: fixed; border-radius: 50%; filter: blur(100px);
    opacity: 0.45; z-index: -1; pointer-events: none;
    animation: drift 16s ease-in-out infinite alternate;
}
.a1 { width: 55vw; height: 55vw; background: var(--c1); top: -20%; left: -12%; }
.a2 { width: 50vw; height: 50vw; background: var(--c2); bottom: -25%; right: -10%; animation-delay: 2s; animation-duration: 20s; }
.a3 { width: 30vw; height: 30vw; background: linear-gradient(var(--c1), var(--c2)); top: 30%; left: 38%; animation-duration: 24s; }
.rcard {
    position: fixed; bottom: -170px; width: 80px; height: 104px;
    border-radius: 8px; border: 2px solid var(--c1);
    background:
        radial-gradient(circle at 25% 22%, var(--c1) 0 8px, transparent 9px),
        linear-gradient(var(--c2), var(--c2)) 45px 12px / 28px 5px no-repeat,
        repeating-linear-gradient(to bottom, transparent 0 8px, var(--border) 8px 11px) 8px 42px / 64px 56px no-repeat;
    opacity: 0; z-index: -1; pointer-events: none;
    animation: riseCard linear infinite;
}
.r1 { left: 4%;  --s: 1.3; animation-duration: 16s; }
.r2 { left: 14%; --s: 0.8; animation-duration: 22s; animation-delay: 3s; }
.r3 { left: 27%; --s: 1.6; animation-duration: 18s; animation-delay: 6s; }
.r4 { left: 45%; --s: 1.0; animation-duration: 24s; animation-delay: 1s; }
.r5 { left: 62%; --s: 1.4; animation-duration: 17s; animation-delay: 8s; }
.r6 { left: 76%; --s: 0.9; animation-duration: 21s; animation-delay: 4s; }
.r7 { left: 90%; --s: 1.2; animation-duration: 19s; animation-delay: 10s; }
@keyframes riseCard {
    0%   { transform: translateY(0) rotate(-10deg) scale(var(--s)); opacity: 0; }
    10%  { opacity: 0.6; }
    100% { transform: translateY(-125vh) rotate(12deg) scale(var(--s)); opacity: 0; }
}

/* Stickers */
.sticker {
    position: fixed; font-size: 42px; opacity: 0.9;
    pointer-events: none; z-index: 1;
    animation: float 4s ease-in-out infinite;
}
.s1 { top: 12%; left: 3%; }
.s2 { top: 25%; right: 4%; animation-delay: 1s; }
.s3 { top: 55%; left: 2%; animation-delay: 2s; }
.s4 { top: 70%; right: 3%; animation-delay: 0.5s; }
.s5 { top: 88%; left: 8%; animation-delay: 1.5s; }
.s6 { top: 8%; right: 12%; animation-delay: 2.5s; }
@keyframes float {
    0%, 100% { transform: translateY(0) rotate(-5deg); }
    50% { transform: translateY(-15px) rotate(5deg); }
}

/* Robot sticker with speech bubble */
.robot-wrap {
    position: fixed; right: 22px; bottom: 16px; z-index: 1;
    display: flex; flex-direction: column; align-items: flex-end;
    pointer-events: none;
}
.robot {
    font-size: 84px; line-height: 1;
    filter: drop-shadow(0 6px 10px rgba(0, 0, 0, 0.35));
    animation: bot 3s ease-in-out infinite;
}
.robot-msg {
    background: var(--solid); color: var(--text);
    border: 1px solid var(--border);
    border-radius: 16px 16px 4px 16px;
    padding: 8px 14px; margin-bottom: 4px;
    font-size: 14px; font-weight: 600;
    box-shadow: 0 4px 14px rgba(0, 0, 0, 0.3);
}
@keyframes bot {
    0%, 100% { transform: translateY(0) rotate(-4deg); }
    50% { transform: translateY(-14px) rotate(4deg); }
}
@media (max-width: 900px) { .robot-wrap { display: none; } }

/* Text */
.stApp h2, .stApp h3, .stApp p, .stApp span, .stApp label, .stApp li {
    color: var(--text) !important;
}
.stApp small { color: var(--text) !important; opacity: 0.8; }
.stApp h1 {
    text-align: center;
    font-weight: 800;
    letter-spacing: 3px;
    background: linear-gradient(90deg, var(--c1), var(--c2));
    -webkit-background-clip: text;
    background-clip: text;
    -webkit-text-fill-color: transparent;
    color: transparent !important;
}

/* Education boxes (expanders) */
[data-testid="stExpander"] {
    border: 1px solid var(--border) !important;
    border-radius: 14px !important;
    background: rgba(255, 255, 255, 0.05) !important;
    margin-bottom: 12px;
}
[data-testid="stExpander"] summary,
[data-testid="stExpander"] summary * {
    color: var(--text) !important;
    font-weight: 700;
}

/* Inputs */
.stTextInput input,
.stTextArea textarea {
    color: var(--text) !important;
    background: var(--input) !important;
    border: 1px solid var(--border) !important;
    border-radius: 10px;
}
.stTextInput input::placeholder,
.stTextArea textarea::placeholder {
    color: var(--text) !important;
    opacity: 0.55;
}
div[data-baseweb="select"] > div {
    background: var(--input) !important;
    border: 1px solid var(--border) !important;
    border-radius: 10px;
}
div[data-baseweb="select"] * { color: var(--text) !important; }
div[data-baseweb="popover"],
div[data-baseweb="popover"] ul { background: var(--solid) !important; }
div[data-baseweb="popover"] li,
div[data-baseweb="popover"] li * { color: var(--text) !important; }

/* File uploader */
[data-testid="stFileUploaderDropzone"] {
    background: var(--input) !important;
    border: 2px dashed var(--c1) !important;
    border-radius: 12px;
}
[data-testid="stFileUploaderDropzone"] button {
    background: var(--card) !important;
    border: 1px solid var(--border) !important;
    border-radius: 10px;
}
[data-testid="stFileUploaderDropzone"] button p,
[data-testid="stFileUploaderDropzone"] button span {
    color: var(--text) !important;
}

/* Tabs */
[data-baseweb="tab-highlight"] { background-color: var(--c1) !important; }
[data-baseweb="tab-border"] { background-color: var(--border) !important; }

/* Buttons */
.stButton > button,
.stDownloadButton > button {
    width: 100%;
    background: linear-gradient(90deg, var(--c1), var(--c2));
    border: none;
    border-radius: 12px;
    padding: 0.6rem;
    font-weight: 700;
    transition: transform 0.2s ease, box-shadow 0.2s ease;
}
.stButton > button p,
.stDownloadButton > button p {
    color: var(--btn) !important;
    font-weight: 700;
}
.stButton > button:hover,
.stDownloadButton > button:hover {
    transform: translateY(-2px);
    box-shadow: 0 6px 22px var(--border);
}

/* Sidebar */
[data-testid="stSidebar"] {
    background: var(--panel) !important;
    backdrop-filter: blur(12px);
    border-right: 1px solid var(--border);
    z-index: 100;
}
[data-testid="stSidebar"] h1,
[data-testid="stSidebar"] h2,
[data-testid="stSidebar"] h3 {
    text-align: left;
    font-size: 1.25rem;
    letter-spacing: 0;
}
[data-testid="stSidebar"] .stButton > button {
    text-align: left;
    justify-content: flex-start;
    background: var(--card);
    border: 1px solid var(--border);
    margin-bottom: 4px;
}
[data-testid="stSidebar"] .stButton > button p {
    color: var(--text) !important;
    font-weight: 600;
}
"""


def bg_layers(style):
    if style == "Animated grid":
        return '<div class="bg-grid"></div><div class="orb o1"></div><div class="orb o2"></div>'
    if style == "Aurora glow":
        return '<div class="blob a1"></div><div class="blob a2"></div><div class="blob a3"></div>'
    if style == "Floating resumes":
        return "".join(f'<div class="rcard r{i}"></div>' for i in range(1, 8))
    return '<div class="bg-dots"></div><div class="orb o1"></div><div class="orb o2"></div>'


photo_html = ""
if CUSTOM_IMAGE_URL:
    photo_html = f"<div class=\"bg-photo\" style=\"background-image:url('{CUSTOM_IMAGE_URL}')\"></div>"

stickers = """
<div class="sticker s1">📄</div>
<div class="sticker s2">⭐</div>
<div class="sticker s3">🚀</div>
<div class="sticker s4">💼</div>
<div class="sticker s5">🎓</div>
<div class="sticker s6">✨</div>
<div class="robot-wrap">
<div class="robot-msg">Hi! I'm your resume bot 👋</div>
<div class="robot">🤖</div>
</div>
"""

full_html = f"<style>{root_css}{main_css}</style>{photo_html}{bg_layers(bg_style)}{stickers}"
full_html = "\n".join(line.strip() for line in full_html.splitlines() if line.strip())
st.markdown(full_html, unsafe_allow_html=True)

# ---------- SIDEBAR: SAVED RESUMES ----------
with st.sidebar:
    st.divider()
    st.header("🕘 Saved Resumes")
    search = st.text_input("Search", placeholder="🔍 Search resumes...")

    shown = 0
    for idx in reversed(range(len(st.session_state.history))):
        item = st.session_state.history[idx]
        if search and search.lower() not in (item["title"] + item["text"]).lower():
            continue
        if st.button("📄 " + item["title"], key=f"hist_{idx}"):
            st.session_state.resume_text = item["text"]
            st.session_state.current = idx
            rerun()
        shown += 1
    if shown == 0:
        st.caption("No resumes yet ✨")

    st.divider()
    if st.button("🗑️ Delete all history"):
        st.session_state.history = []
        st.session_state.resume_text = ""
        st.session_state.current = None
        rerun()


# ---------- FILE HELPERS ----------
def describe_files(files):
    """Turns uploaded files into text for the AI (file name + text inside PDFs)."""
    if not files:
        return "None"
    out = []
    for f in files:
        info = f"- File: {f.name}"
        if f.name.lower().endswith(".pdf") and PdfReader:
            try:
                reader = PdfReader(io.BytesIO(f.getvalue()))
                text = " ".join((p.extract_text() or "") for p in reader.pages[:2])
                text = " ".join(text.split())[:800]
                if text:
                    info += f"\n  Text found in file: {text}"
            except Exception:
                pass
        out.append(info)
    return "\n".join(out)


# ---------- EDUCATION LINE BUILDERS ----------
def clean(v):
    """Turns '-- Select --' and empty values into an empty string."""
    v = (v or "").strip()
    return "" if v == PH else v


def score_line(kind, value):
    value = clean(value)
    return f"{kind}: {value}" if value else ""


def join_line(parts):
    return " | ".join(p for p in parts if p)


def degree_line(f):
    degree = clean(f["degree"])
    if degree == "Other" and clean(f["degree_other"]):
        degree = clean(f["degree_other"])
    branch = clean(f["branch"])
    if branch == "Other" and clean(f["branch_other"]):
        branch = clean(f["branch_other"])
    college = clean(f["college"])
    if not any([degree, branch, college]):
        return ""
    if degree and branch:
        title = f"{degree} in {branch}"
    else:
        title = degree or branch or "Degree"
    place = ", ".join(p for p in [college, clean(f["city"])] if p)
    start, end = clean(f["start"]), clean(f["end"])
    if f["current"]:
        end = f"Expected {end}" if end else "Present"
    years = f"{start or '?'} - {end or '?'}" if (start or end) else ""
    return join_line([
        f"**{title}**", place, clean(f["univ"]), years,
        score_line(f["stype"], f["score"]),
    ])


def inter_line(f):
    college, stream = clean(f["college"]), clean(f["stream"])
    if not (college or stream):
        return ""
    title = "Intermediate / 12th" + (f" ({stream})" if stream else "")
    start, end = clean(f["start"]), clean(f["end"])
    years = f"{start or '?'} - {end or '?'}" if (start or end) else ""
    return join_line([
        f"**{title}**", college, clean(f["board"]), years,
        score_line(f["stype"], f["score"]),
    ])


def tenth_line(f):
    school = clean(f["school"])
    if not school:
        return ""
    return join_line([
        "**10th / SSC**", school, clean(f["board"]), clean(f["year"]),
        score_line(f["stype"], f["score"]),
    ])


def show_line(line):
    if line:
        st.markdown("✅ **Will appear as:**\n\n- " + line)
    else:
        st.caption("Fill the boxes above to see how this will appear.")


# ---------- EDUCATION BOXES (one per qualification) ----------
def degree_box(prefix, title, expanded):
    with st.expander(title, expanded=expanded):
        a, b = st.columns(2)
        degree = a.selectbox("Degree", DEGREES, key=f"{prefix}_degree")
        branch = b.selectbox("Branch", BRANCHES, key=f"{prefix}_branch")
        o1, o2 = st.columns(2)
        degree_other = o1.text_input("If degree is 'Other', type it here", key=f"{prefix}_degree_other")
        branch_other = o2.text_input("If branch is 'Other', type it here", key=f"{prefix}_branch_other")
        c, d = st.columns(2)
        college = c.text_input("College name", key=f"{prefix}_college")
        city = d.text_input("College city", key=f"{prefix}_city")
        univ = st.text_input("University (optional)", key=f"{prefix}_univ")
        e, g = st.columns(2)
        start = e.selectbox("Start year", YEARS, key=f"{prefix}_start")
        end = g.selectbox("End year / expected", YEARS, key=f"{prefix}_end")
        current = st.checkbox("I am currently studying here", key=f"{prefix}_current")
        h, i = st.columns(2)
        stype = h.radio("Score type", ["CGPA", "Percentage"], horizontal=True, key=f"{prefix}_stype")
        score = i.text_input("Score", placeholder="e.g. 8.2 or 78%", key=f"{prefix}_score")
        line = degree_line({
            "degree": degree, "degree_other": degree_other, "branch": branch,
            "branch_other": branch_other, "college": college, "city": city,
            "univ": univ, "start": start, "end": end, "current": current,
            "stype": stype, "score": score,
        })
        show_line(line)
    return line


def inter_box():
    with st.expander("📚 Intermediate / 12th", expanded=True):
        a, b = st.columns(2)
        college = a.text_input("College name", key="inter_college")
        stream = b.selectbox("Stream / Group", STREAMS, key="inter_stream")
        c, d, e = st.columns(3)
        board = c.selectbox("Board", BOARDS_12, key="inter_board")
        start = d.selectbox("Start year", YEARS, key="inter_start")
        end = e.selectbox("End year", YEARS, key="inter_end")
        g, h = st.columns(2)
        stype = g.radio("Score type", ["Percentage", "CGPA"], horizontal=True, key="inter_stype")
        score = h.text_input("Score", placeholder="e.g. 95% or 9.4", key="inter_score")
        line = inter_line({
            "college": college, "stream": stream, "board": board,
            "start": start, "end": end, "stype": stype, "score": score,
        })
        show_line(line)
    return line


def tenth_box():
    with st.expander("🏫 10th / SSC", expanded=True):
        a, b = st.columns(2)
        school = a.text_input("School name", key="tenth_school")
        board = b.selectbox("Board", BOARDS_10, key="tenth_board")
        c, d, e = st.columns(3)
        year = c.selectbox("Year of passing", YEARS, key="tenth_year")
        stype = d.radio("Score type", ["CGPA", "Percentage"], horizontal=True, key="tenth_stype")
        score = e.text_input("Score", placeholder="e.g. 9.5 or 92%", key="tenth_score")
        line = tenth_line({
            "school": school, "board": board, "year": year,
            "stype": stype, "score": score,
        })
        show_line(line)
    return line


def add_education(resume, entries):
    """Removes any Education section written by the AI and inserts ours."""
    if not entries:
        return resume
    edu_md = "## Education\n\n" + "\n".join(f"- {e}" for e in entries)

    resume = re.sub(
        r"(?ims)^#{1,3}\s*education\b.*?(?=^#{1,2}\s|\Z)", "", resume
    ).rstrip()

    m = re.search(r"(?im)^#{1,3}\s*(certifications?|achievements?)\b", resume)
    if m:
        return resume[: m.start()] + edu_md + "\n\n" + resume[m.start():]
    return resume + "\n\n" + edu_md


# ---------- PROMPT + GENERATION ----------
def build_prompt(d):
    return f"""Create a complete, professional resume in Markdown for the person below.

STRICT RULES:
- Use ONLY the information provided. Do NOT invent companies, degrees, dates, numbers or skills.
- If a section has no information, leave that section out.
- Do NOT write an Education section. It is added separately by the app.
- Certifications and Achievements: use only the typed text and the uploaded file details below.
  If an uploaded file has no readable text, use a cleaned-up version of its file name as the title only.
  Never guess issuers, dates or scores.
- Write strong bullet points starting with action verbs.
- Style: {d['template']}. Tone: {d['tone']}.
- {"Make it ATS-friendly: simple headings, no tables, no icons." if d['ats'] else "Normal formatting is fine."}
- Use "# Name" for the name, "## Section" for section headings, and "-" for bullets.
- Sections in order: Name and contact line, Professional Summary, Skills, Experience, Projects, Certifications, Achievements.
- Output only the resume. No introduction and no closing comments.

PERSON DETAILS:
Name: {d['name']}
Target job role: {d['role']}
Email: {d['email']}
Phone: {d['phone']}
Location: {d['location']}
LinkedIn: {d['linkedin']}
GitHub / Portfolio: {d['github']}

Highest qualification (for the summary only): {d['edu_brief']}

Skills:
{d['skills']}

Experience / Internships:
{d['experience']}

Projects:
{d['projects']}

Certifications (typed):
{d['certs']}

Certificates (uploaded files):
{d['cert_files']}

Achievements (typed):
{d['achievements']}

Achievements (uploaded files):
{d['ach_files']}

Target job description (tailor the resume to this if given):
{d['jd']}
"""


def generate(prompt):
    try:
        stream = client.chat(
            model=model_name,
            messages=[
                {"role": "system", "content": "You are an expert resume writer."},
                {"role": "user", "content": prompt},
            ],
            stream=True,
        )
        for chunk in stream:
            yield chunk["message"]["content"]
    except Exception as e:
        yield f"\n\n⚠️ Could not reach Ollama: {e}"


def stream_text(gen):
    """Shows text as it is written (works on every Streamlit version)."""
    box = st.empty()
    acc = ""
    for chunk in gen:
        acc += chunk
        box.markdown(acc + "▌")
    box.markdown(acc)
    return acc


# ---------- DOWNLOAD BUILDERS (PDF / WORD) ----------
def strip_md(text):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"(?<!\*)\*(?!\*)", "", text)
    return text.strip()


def keep_bold(text):
    """Keeps **bold** markers, removes other markdown symbols."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("__", "").replace("`", "")
    text = re.sub(r"(?<!\*)\*(?!\*)", "", text)
    return text.strip()


def parse_md(md):
    items = []
    for raw in md.splitlines():
        s = raw.strip()
        if not s:
            items.append(("space", ""))
        elif len(s) >= 3 and set(s) <= set("-_*= "):
            continue
        elif s.startswith("### "):
            items.append(("h3", s[4:]))
        elif s.startswith("## "):
            items.append(("h2", s[3:]))
        elif s.startswith("# "):
            items.append(("h1", s[2:]))
        elif s.startswith(("- ", "* ", "• ")):
            items.append(("bullet", s[2:]))
        else:
            items.append(("p", s))
    return items


def make_docx(md):
    if Document is None:
        return None
    try:
        doc = Document()
        levels = {"h1": 0, "h2": 1, "h3": 2}
        for kind, text in parse_md(md):
            if kind == "space":
                continue
            if kind in levels:
                doc.add_heading(strip_md(text), level=levels[kind])
                continue
            para = doc.add_paragraph(style="List Bullet") if kind == "bullet" else doc.add_paragraph()
            for i, part in enumerate(keep_bold(text).split("**")):
                run = para.add_run(part)
                run.bold = (i % 2 == 1)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    except Exception:
        return None


def latin(text):
    swaps = {"•": "-", "–": "-", "—": "-", "“": '"', "”": '"', "‘": "'", "’": "'",
             "…": "...", "\u00a0": " ", "--": "-"}
    for k, v in swaps.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def make_pdf(md):
    if FPDF is None:
        return None
    try:
        pdf = FPDF()
        pdf.set_margins(15, 15, 15)
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        for kind, text in parse_md(md):
            text = latin(text)
            if kind == "space":
                pdf.ln(2)
            elif kind == "h1":
                pdf.set_font("Helvetica", "B", 20)
                pdf.multi_cell(0, 10, strip_md(text), align="C", new_x="LMARGIN", new_y="NEXT")
            elif kind == "h2":
                pdf.ln(3)
                pdf.set_font("Helvetica", "B", 13)
                pdf.multi_cell(0, 7, strip_md(text).upper(), new_x="LMARGIN", new_y="NEXT")
                y = pdf.get_y()
                pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
                pdf.ln(1)
            elif kind == "h3":
                pdf.set_font("Helvetica", "B", 11)
                pdf.multi_cell(0, 6, strip_md(text), new_x="LMARGIN", new_y="NEXT")
            elif kind == "bullet":
                pdf.set_font("Helvetica", "", 10.5)
                pdf.set_x(pdf.l_margin + 4)
                pdf.multi_cell(0, 5.5, "- " + keep_bold(text), markdown=True, new_x="LMARGIN", new_y="NEXT")
            else:
                pdf.set_font("Helvetica", "", 10.5)
                pdf.multi_cell(0, 5.5, keep_bold(text), markdown=True, new_x="LMARGIN", new_y="NEXT")
        return bytes(pdf.output())
    except Exception:
        return None


# ---------- MAIN PAGE ----------
st.title("AI RESUME BUILDER")
st.caption("Fill in your details and let AI write your resume ✨")

st.subheader("👤 Personal details")
c1, c2 = st.columns(2)
name = c1.text_input("Full name *")
role = c2.text_input("Target job role *", placeholder="e.g. Python Developer")
c3, c4 = st.columns(2)
email = c3.text_input("Email")
phone = c4.text_input("Phone")
c5, c6, c7 = st.columns(3)
location = c5.text_input("Location", placeholder="City, State")
linkedin = c6.text_input("LinkedIn link")
github = c7.text_input("GitHub / Portfolio link")

# ---------- EDUCATION ----------
st.subheader("🎓 Education")
st.caption("Fill each box. The line under it shows exactly how it will look in your resume.")

pg_line = degree_box("pg", "🎓 Post-graduation / Other degree (optional, e.g. M.Tech, MCA, MBA)", False)
deg_line = degree_box("deg", "🎓 Degree (B.Tech / B.E / B.Sc / etc.)", True)
int_line = inter_box()
ten_line = tenth_box()

entries = [e for e in [pg_line, deg_line, int_line, ten_line] if e]

if entries:
    st.markdown("### 📋 Education preview")
    st.markdown("\n".join(f"- {e}" for e in entries))

# ---------- SKILLS, CERTIFICATES, ACHIEVEMENTS ----------
st.subheader("🛠️ Skills, experience and projects")
skills = st.text_area("Skills", placeholder="Python, Streamlit, SQL, Machine Learning", height=90)
experience = st.text_area("Experience / Internships", placeholder="Company, role, duration, what you did", height=110)
projects = st.text_area("Projects", placeholder="Project name, tools used, what it does", height=110)

st.subheader("🏅 Certificates")
certs = st.text_area("Certifications (type names, optional)", height=80)
cert_files = st.file_uploader(
    "📎 Upload your certificates",
    type=["pdf", "png", "jpg", "jpeg"],
    accept_multiple_files=True,
    key="cert_uploader",
)

st.subheader("🏆 Achievements")
achievements = st.text_area("Achievements (type them, optional)", height=80)
ach_files = st.file_uploader(
    "📎 Upload achievement proofs",
    type=["pdf", "png", "jpg", "jpeg"],
    accept_multiple_files=True,
    key="ach_uploader",
)

st.subheader("🎯 Customize")
jd = st.text_area("Job description (optional, paste to tailor the resume)", height=100)
o1, o2, o3 = st.columns(3)
template = o1.selectbox("Style", ["Fresher / Student", "Professional", "Modern"])
tone = o2.selectbox("Tone", ["Confident", "Formal", "Friendly"])
ats = o3.checkbox("ATS-friendly", value=True)

generate_clicked = st.button("✨ Generate Resume")

if generate_clicked:
    if not name.strip() or not role.strip():
        st.warning("Please fill in your name and target job role.")
    else:
        edu_brief = strip_md(entries[0]) if entries else "None"
        data = {
            "name": name, "role": role, "email": email, "phone": phone,
            "location": location, "linkedin": linkedin, "github": github,
            "edu_brief": edu_brief, "skills": skills, "experience": experience,
            "projects": projects, "certs": certs, "achievements": achievements,
            "cert_files": describe_files(cert_files),
            "ach_files": describe_files(ach_files),
            "jd": jd, "template": template, "tone": tone, "ats": ats,
        }

        st.subheader("📄 Writing your resume...")
        text = stream_text(generate(build_prompt(data)))

        if text and "Could not reach Ollama" not in text:
            final_resume = add_education(text, entries)
            st.session_state.history.append({"title": f"{name} - {role}", "text": final_resume})
            st.session_state.current = len(st.session_state.history) - 1
            st.session_state.resume_text = final_resume
            rerun()

# ---------- RESULT: PREVIEW / EDIT / DOWNLOAD ----------
if st.session_state.resume_text:
    st.subheader("📄 Your Resume")
    tab_preview, tab_edit = st.tabs(["👁️ Preview", "✏️ Edit"])

    with tab_preview:
        st.markdown(st.session_state.resume_text)

    with tab_edit:
        st.caption("Change anything you like, then press Ctrl+Enter or click outside the box to update the preview.")
        st.text_area("Edit your resume", key="resume_text", height=520)
        if st.button("💾 Save changes to history"):
            idx = st.session_state.current
            if idx is not None and idx < len(st.session_state.history):
                st.session_state.history[idx]["text"] = st.session_state.resume_text
                st.success("Saved ✅")

    st.subheader("⬇️ Download")
    resume_md = st.session_state.resume_text
    pdf_bytes = make_pdf(resume_md)
    docx_bytes = make_docx(resume_md)

    d1, d2, d3 = st.columns(3)
    if pdf_bytes:
        d1.download_button("📕 PDF", pdf_bytes, file_name="resume.pdf", mime="application/pdf")
    else:
        d1.caption("PDF needs `fpdf2` in requirements.txt")
    if docx_bytes:
        d2.download_button(
            "📘 Word (.docx)", docx_bytes, file_name="resume.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    else:
        d2.caption("Word needs `python-docx` in requirements.txt")
    d3.download_button("📄 Text (.txt)", resume_md, file_name="resume.txt", mime="text/plain")